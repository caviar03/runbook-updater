import os
from pathlib import Path
from tempfile import NamedTemporaryFile

from docx import Document
from docx.oxml.ns import qn
from docx.text.run import Run

from src.models.change_proposal import changeProposal


def _paragraphs(container):
    for paragraph in container.paragraphs:
        yield paragraph
    visited_cells = set()
    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                # Merged cells may appear more than once in the row/column grid.
                if cell._tc in visited_cells:
                    continue
                visited_cells.add(cell._tc)
                yield from _paragraphs(cell)


def _runs(paragraph):
    for item in paragraph.iter_inner_content():
        if isinstance(item, Run):
            yield item
        else:
            yield from item.runs


def _replace(paragraph, current_text, proposed_text):
    runs = list(_runs(paragraph))
    text = "".join(run.text for run in runs)
    if text != paragraph.text:
        raise ValueError("This paragraph contains unsupported text structures")
    start = text.index(current_text)
    end = start + len(current_text)
    positions = []
    cursor = 0
    for i, run in enumerate(runs):
        positions.append((i, cursor, cursor + len(run.text)))
        cursor += len(run.text)
    first, first_start, _ = next(item for item in positions if item[1] <= start < item[2])
    last, last_start, _ = next(item for item in positions if item[1] < end <= item[2])
    supported = {qn(f"w:{tag}") for tag in ["rPr", "t", "tab", "br", "cr"]}
    for run in runs[first:last + 1]:
        if any(element.tag not in supported for element in run._r):
            raise ValueError("Approved text overlaps an image, field, or unsupported run content")
    prefix = runs[first].text[:start - first_start]
    suffix = runs[last].text[end - last_start:]
    if first == last:
        runs[first].text = prefix + proposed_text + suffix
    else:
        runs[first].text = prefix + proposed_text
        for i in range(first + 1, last):
            runs[i].text = ""
        runs[last].text = suffix


def prepare_document(runbook_path: str, approved_changes: changeProposal):
    if not approved_changes or not approved_changes.changes:
        raise ValueError("Nonempty human-approved changes are required")
    document = Document(runbook_path)
    paragraphs = list(_paragraphs(document))
    for change in approved_changes.changes:
        if not change.current_text or not change.current_text.strip():
            raise ValueError("An update must identify exact nonempty current text")
        matches = [paragraph for paragraph in paragraphs if change.current_text in paragraph.text]
        occurrences = sum(paragraph.text.count(change.current_text) for paragraph in matches)
        if occurrences != 1:
            raise ValueError(f"Approved text for {change.section!r} must match exactly once; found {occurrences}")
        _replace(matches[0], change.current_text, change.proposed_text)
    return document


def read_runbook_text(runbook_path: str) -> str:
    return "\n\n".join(paragraph.text for paragraph in _paragraphs(Document(runbook_path)))


def implementor(runbook_path: str, approved_changes: changeProposal, output_path: str) -> str:
    document = prepare_document(runbook_path, approved_changes)
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    # Publish only after every approved edit succeeds; failures leave existing files intact.
    temporary = None
    try:
        with NamedTemporaryFile(dir=output.parent, suffix=".docx", delete=False) as file:
            temporary = Path(file.name)
        document.save(temporary)
        os.replace(temporary, output)
    finally:
        if temporary is not None:
            temporary.unlink(missing_ok=True)
    return str(output)
