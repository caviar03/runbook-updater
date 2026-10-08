import sys
from pathlib import Path
from types import ModuleType, SimpleNamespace
from unittest.mock import Mock

import pytest
from docx import Document

from src.agents.implementor import implementor, read_runbook_text
from src.graph.state import make_initial_state
from src.graph.workflow import build_workflow
from src.models.change_proposal import changeProposal, proposedChange
from test_workflow import setup_workflow


def doc(path, text="Wait 30 seconds"):
    document = Document()
    document.add_paragraph("Recovery", style="Heading 1")
    document.add_paragraph(text)
    document.save(path)
    return path


def change(runbook, old="Wait 30 seconds", new="Wait 60 seconds", section="Recovery"):
    return proposedChange(runbook=runbook, section=section, current_text=old,
                          proposed_text=new, reason="The configured timeout changed")


def test_multi_runbook_graph_updates_only_the_named_files_after_approval(monkeypatch, tmp_path):
    state, _, _, calls = setup_workflow(monkeypatch, tmp_path)
    first = doc(tmp_path / "Orders.docx")
    second = doc(tmp_path / "Payments.docx")
    untouched = doc(tmp_path / "Search.docx")
    output = tmp_path / "output"
    state = make_initial_state({key: state[key] for key in ("old_code", "current_code", "code_diff", "changed_files")},
                               [first, second, untouched], output)
    calls.proposer.return_value = changeProposal(readable_feedback="Two runbooks need an update",
                                                changes=[change(first.name), change(second.name)])
    monkeypatch.setattr("builtins.input", lambda _: "yes")
    result = build_workflow().invoke(state)
    assert result["workflow_status"] == "applied"
    assert set(result["updated_runbook_paths"]) == {str(output / first.name), str(output / second.name)}
    assert Document(output / first.name).paragraphs[1].text == "Wait 60 seconds"
    assert Document(output / second.name).paragraphs[1].text == "Wait 60 seconds"
    assert not (output / untouched.name).exists()
    assert all(Document(path).paragraphs[1].text == "Wait 30 seconds" for path in (first, second, untouched))


def test_multi_runbook_writer_rejects_unknown_destinations_and_incomplete_edits_before_saving(tmp_path):
    first, second = doc(tmp_path / "Orders.docx"), doc(tmp_path / "Payments.docx")
    for bad_change in (change("Unknown.docx"), change(second.name, old="Missing text")):
        proposal = changeProposal(readable_feedback="Review", changes=[change(first.name), bad_change])
        with pytest.raises(ValueError):
            implementor(runbook_paths=[first, second], approved_changes=proposal, output_dir=tmp_path / "output")
        assert not (tmp_path / "output").exists()


def test_section_disambiguates_repeated_text_without_losing_other_content(tmp_path):
    source, output = tmp_path / "Orders.docx", tmp_path / "updated.docx"
    document = Document()
    for section in ("Recovery", "Deployment"):
        document.add_paragraph(section, style="Heading 1")
        document.add_paragraph("Wait 30 seconds")
    document.save(source)
    proposal = changeProposal(readable_feedback="Review", changes=[change(source.name)])
    implementor(str(source), proposal, str(output))
    assert [paragraph.text for paragraph in Document(output).paragraphs] == ["Recovery", "Wait 60 seconds", "Deployment", "Wait 30 seconds"]


def test_flattened_table_changes_require_all_cells_to_apply(tmp_path):
    source, output = tmp_path / "Orders.docx", tmp_path / "updated.docx"
    document = Document()
    document.add_paragraph("Recovery", style="Heading 1")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Wait 30 seconds"; table.cell(0, 1).text = "2 retries"
    document.save(source)
    old = "Retry, **Timeout** = Wait 30 seconds. Retry, **Count** = 2 retries."
    new = "Retry, **Timeout** = Wait 60 seconds. Retry, **Count** = 3 retries."
    proposal = changeProposal(readable_feedback="Review", changes=[change(source.name, old, new)])
    implementor(str(source), proposal, str(output))
    assert Document(output).tables[0].cell(0, 0).text == "Wait 60 seconds"
    assert Document(output).tables[0].cell(0, 1).text == "3 retries"
    before = output.read_bytes()
    proposal.changes[0].current_text = old.replace("2 retries", "missing cell")
    with pytest.raises(ValueError):
        implementor(str(source), proposal, str(output))
    assert output.read_bytes() == before


def test_direct_docx_context_preserves_runbook_identity_and_nested_section_labels(tmp_path):
    source = doc(tmp_path / "Orders.docx")
    document = Document(source)
    document.add_paragraph("Verification", style="Heading 2")
    cell = document.add_table(rows=1, cols=1).cell(0, 0)
    cell.add_table(rows=1, cols=1).cell(0, 0).text = "Check retries"
    document.save(source)
    context = read_runbook_text(str(source), runbook_name=source.name)
    assert "[Runbook: Orders.docx]" in context
    assert "[Section: Recovery > Verification]\nCheck retries" in context


def test_cli_indexes_all_selected_runbooks_before_invoking_the_graph(monkeypatch, tmp_path):
    import src.main as main
    folder = tmp_path / "test_data/input"; folder.mkdir(parents=True)
    paths = [doc(folder / name) for name in ("Orders.docx", "Payments.docx")]
    monkeypatch.chdir(tmp_path)
    for key, value in {"GITHUB_OWNER": "owner", "GITHUB_REPO": "code", "GITHUB_PR_NUMBER": "3", "GITHUB_MERGE_COMMIT_SHA": "a" * 40}.items():
        monkeypatch.setenv(key, value)
    monkeypatch.delenv("RUNBOOK_PATHS", raising=False)
    monkeypatch.delenv("RUNBOOK_OUTPUT_DIR", raising=False)
    indexer = Mock()
    fake_index_module = ModuleType("src.retrieval.index_runbooks")
    fake_index_module.index_runbooks = indexer
    monkeypatch.setitem(sys.modules, "src.retrieval.index_runbooks", fake_index_module)
    monkeypatch.setattr(main, "get_merge_code", lambda *args, **kwargs: {"old_code": "old", "current_code": "new", "code_diff": "diff", "changed_files": []})
    invoke = Mock(return_value={"workflow_status": "no_update_needed"})
    monkeypatch.setattr(main, "build_workflow", lambda: SimpleNamespace(invoke=invoke))
    main.main()
    selected = [str(path.relative_to(tmp_path)) for path in paths]
    indexer.assert_called_once_with(selected)
    assert invoke.call_args.args[0]["runbook_paths"] == selected


def test_legacy_proposal_without_runbook_id_still_parses_and_applies(tmp_path):
    source, output = doc(tmp_path / "Orders.docx"), tmp_path / "updated.docx"
    proposal = changeProposal.model_validate({"readable_feedback": "Saved before merge", "approved_changes": None,
        "changes": [{"section": "Recovery", "current_text": "Wait 30 seconds", "proposed_text": "Wait 60 seconds", "reason": "Timeout changed"}]})
    implementor(str(source), proposal, str(output))
    assert Document(output).paragraphs[1].text == "Wait 60 seconds"
