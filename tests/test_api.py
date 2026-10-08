import json
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from unittest.mock import Mock

import pytest
from docx import Document
from fastapi.testclient import TestClient

from src.api.app import create_app
from src.api.models import StartAnalysis
from src.api.service import ApiError, Service
from src.api.settings import Settings
from test_workflow import setup_workflow


TOKEN = "test-service-token-with-32-characters"
AUTH = {"Authorization": f"Bearer {TOKEN}"}


def request():
    return StartAnalysis.model_validate({
        "request_id": "github-run-1", "attempt": 1,
        "source": {"repository": "owner/code", "pull_request_number": 42,
                   "pull_request_url": "https://github.com/owner/code/pull/42",
                   "base_sha": "a" * 40, "head_sha": "b" * 40, "merge_commit_sha": "c" * 40,
                   "merged_at": "2026-10-08T01:00:00Z"},
        "trigger": {"type": "pull_request_merged", "actor": "developer"},
        "code_snapshot": {"head_sha": "b" * 40, "merge_commit_sha": "c" * 40, "files": [
            {"file_path": "worker.py", "previous_file_path": "worker.py", "status": "modified",
             "additions": 1, "deletions": 1, "patch": "@@ -1 +1 @@\n-timeout = 30\n+timeout = 60\n",
             "pr_head_content": "timeout = 60\n", "merged_content": "timeout = 60\n"}
        ]},
    }).model_dump(mode="json")


def decision(snapshot, action="approve", actor="reviewer"):
    return {"decision": action, "proposal_version": snapshot["proposal_version"],
            "head_sha": snapshot["source_head_sha"], "merge_commit_sha": snapshot["source_merge_commit_sha"], "actor": actor}


@pytest.fixture
def backend(monkeypatch, tmp_path):
    state, source, _, calls = setup_workflow(monkeypatch, tmp_path)
    settings = Settings(TOKEN, "owner/code", source, tmp_path / "state")
    monkeypatch.delenv("GITHUB_TOKEN", raising=False)
    monkeypatch.setattr("builtins.input", Mock(side_effect=AssertionError("API must never ask for terminal input")))
    # Both full-content analyzers and proposal generation must bypass retrieval.
    for name in ("runbook_analyzer1", "runbook_analyzer2", "change_proposer"):
        monkeypatch.setattr(f"src.agents.{name}.retrieve_for_queries", Mock(side_effect=AssertionError("Retrieval is disabled")))
    return Service(settings), source, calls


def review(service):
    status, queued = service.start_analysis(request(), "start-1")
    assert status == 202 and queued["status"] == "queued"
    assert service.run_once()
    result = service.get_analysis(queued["analysis_id"])
    assert result["status"] == "awaiting_decision"
    return result


def test_http_endpoints_auth_validation_and_no_work_during_post(backend):
    service, _, calls = backend
    with TestClient(create_app(service.settings, service, start_worker=False)) as client:
        assert client.get("/health").json()["retrieval"] == "disabled"
        assert client.get("/v1/analyses/unknown").status_code == 401
        assert client.post("/v1/analyses", json=request()).status_code == 401
        assert client.get("/v1/analyses/unknown", headers=AUTH).status_code == 404
        assert client.post("/v1/analyses", json=request(), headers=AUTH).status_code == 422
        body = request(); body["trigger"]["type"] = "pull_request_approved"
        assert client.post("/v1/analyses", json=body, headers=AUTH | {"Idempotency-Key": "invalid"}).status_code == 422
        body = request(); body["source"]["repository"] = "other/repo"; body["source"]["pull_request_url"] = "https://github.com/other/repo/pull/42"
        assert client.post("/v1/analyses", json=body, headers=AUTH | {"Idempotency-Key": "disabled"}).status_code == 403
        response = client.post("/v1/analyses", json=request(), headers=AUTH | {"Idempotency-Key": "start-1"})
        assert response.status_code == 202 and calls.reviewer.call_count == 0
        analysis_id = response.json()["analysis_id"]
        service.run_once()
        snapshot = client.get(f"/v1/analyses/{analysis_id}", headers=AUTH).json()
        response = client.post(f"/v1/analyses/{analysis_id}/decisions", json=decision(snapshot), headers=AUTH | {"Idempotency-Key": "approve-1"})
        assert response.status_code == 202
        assert response.json()["accepted"] is True
        assert response.json()["analysis"]["status"] == "applying"


def test_review_is_saved_without_write_and_approval_survives_restart(backend):
    service, source, calls = backend
    before = source.read_bytes()
    snapshot = review(service)
    target = snapshot["proposals"][0]["target"]
    assert target["kind"] == "local" and target["path"] == source.name
    assert len(target["revision_sha256"]) == 64
    output = service.settings.data_dir / snapshot["analysis_id"] / "updated.docx"
    assert not output.exists() and source.read_bytes() == before
    for analyzer in (calls.analyzer1, calls.analyzer2, calls.proposer):
        assert "Wait 30 seconds" in analyzer.call_args.kwargs["runbook_text"]
        assert f"[Runbook: {source.name}]" in analyzer.call_args.kwargs["runbook_text"]
    restarted = Service(service.settings)
    assert restarted.get_analysis(snapshot["analysis_id"]) == snapshot
    assert not restarted.run_once()  # Awaiting review never re-runs agents on restart.
    receipt = restarted.decide(snapshot["analysis_id"], decision(snapshot), "approve-1")
    assert receipt[1]["analysis"]["status"] == "applying" and not output.exists()
    restarted = Service(service.settings)
    assert restarted.run_once()
    assert restarted.get_analysis(snapshot["analysis_id"])["status"] == "applied"
    assert Document(output).paragraphs[0].text == "Wait 60 seconds"
    assert source.read_bytes() == before
    assert calls.reviewer.call_count == 1
    assert restarted.decide(snapshot["analysis_id"], decision(snapshot), "approve-1") == receipt
    assert not restarted.run_once()


def test_start_receipts_replay_original_response_after_progress_and_restart(backend):
    service, _, calls = backend
    receipt = service.start_analysis(request(), "start-1")
    service.run_once()
    restarted = Service(service.settings)
    assert restarted.start_analysis(request(), "start-1") == receipt
    assert not restarted.run_once() and calls.reviewer.call_count == 1
    changed = request(); changed["attempt"] = 2
    with pytest.raises(ApiError) as error:
        restarted.start_analysis(changed, "start-1")
    assert error.value.status == 409
    with pytest.raises(ApiError) as error:
        restarted.start_analysis(request(), "different-key")
    assert error.value.status == 409
    assert restarted.start_analysis(changed, "start-2")[1]["analysis_id"] != receipt[1]["analysis_id"]


@pytest.mark.parametrize("field,value", [("head_sha", "d" * 40), ("merge_commit_sha", "e" * 40), ("proposal_version", "obsolete")])
def test_stale_decision_cannot_write(backend, field, value):
    service, _, _ = backend
    snapshot = review(service)
    body = decision(snapshot); body[field] = value
    with pytest.raises(ApiError) as error:
        service.decide(snapshot["analysis_id"], body, "stale")
    assert error.value.status == 409
    assert service.get_analysis(snapshot["analysis_id"])["status"] == "awaiting_decision"
    assert not service.run_once()


@pytest.mark.parametrize("modify", ["source", "snapshot", "delete"])
def test_changed_document_is_stale_before_acceptance(backend, modify):
    service, source, _ = backend
    snapshot = review(service)
    if modify == "delete":
        source.unlink()
    else:
        path = source if modify == "source" else Path(service.store.get(snapshot["analysis_id"])["execution"]["state"]["runbook_path"])
        document = Document(); document.add_paragraph("Someone changed this"); document.save(path)
    status, receipt = service.decide(snapshot["analysis_id"], decision(snapshot), "stale-doc")
    assert status == 409 and "accepted" not in receipt
    assert service.get_analysis(snapshot["analysis_id"])["status"] == "stale"
    assert service.store.get(snapshot["analysis_id"])["decision"] is None
    assert not service.run_once()


def test_denial_is_durable_and_never_writes(backend):
    service, source, _ = backend
    before = source.read_bytes(); snapshot = review(service)
    body = decision(snapshot, "deny")
    receipt = service.decide(snapshot["analysis_id"], body, "deny-1")
    assert receipt[0] == 200 and receipt[1]["analysis"]["status"] == "rejected"
    restarted = Service(service.settings)
    assert not restarted.run_once() and source.read_bytes() == before
    assert restarted.decide(snapshot["analysis_id"], body, "deny-1") == receipt
    with pytest.raises(ApiError) as error:
        restarted.decide(snapshot["analysis_id"], decision(snapshot), "approve-late")
    assert error.value.status == 409


def test_no_update_requires_acknowledgment_and_saves_actor(backend):
    from test_workflow import review as runbook_review
    service, _, calls = backend
    calls.judge.return_value = runbook_review(False)
    queued = service.start_analysis(request(), "start-1")[1]; service.run_once()
    snapshot = service.get_analysis(queued["analysis_id"])
    assert snapshot["status"] == "no_update_needed"
    assert snapshot["allowed_decisions"] == ["acknowledge"] and snapshot["proposal_version"]
    with pytest.raises(ApiError) as error:
        service.decide(snapshot["analysis_id"], decision(snapshot), "invalid-approve")
    assert error.value.status == 422
    body = decision(snapshot, "acknowledge", actor="any-github-user")
    receipt = service.decide(snapshot["analysis_id"], body, "ack-1")
    restarted = Service(service.settings)
    assert restarted.decide(snapshot["analysis_id"], body, "ack-1") == receipt
    assert restarted.store.get(snapshot["analysis_id"])["decision"]["actor"] == "any-github-user"
    assert not restarted.run_once()


def test_create_not_offered_without_backend_support(backend):
    service, _, _ = backend
    snapshot = review(service)
    assert snapshot["allowed_decisions"] == ["approve", "deny"]
    with pytest.raises(ApiError) as error:
        service.decide(snapshot["analysis_id"], decision(snapshot, "create"), "create-1")
    assert error.value.status == 422


def test_competing_decisions_have_one_winner_and_conflicting_key_returns_409(backend):
    service, _, _ = backend
    snapshot = review(service)
    def submit(action):
        try:
            return service.decide(snapshot["analysis_id"], decision(snapshot, action), action)[0]
        except ApiError as error:
            return error.status
    with ThreadPoolExecutor(2) as executor:
        results = list(executor.map(submit, ["approve", "deny"]))
    assert sorted(results) == [200 if service.store.get(snapshot["analysis_id"])["decision"]["decision"] == "deny" else 202, 409]
    winning = service.store.get(snapshot["analysis_id"])["decision"]["decision"]
    with pytest.raises(ApiError) as error:
        service.decide(snapshot["analysis_id"], decision(snapshot, "deny" if winning == "approve" else "approve"), winning)
    assert error.value.status == 409


def test_queued_jobs_and_prepared_analysis_resume_after_restart(backend):
    service, _, calls = backend
    queued = service.start_analysis(request(), "start-1")[1]
    # Simulate exit after code/document inputs are saved but before agents finish.
    job = service.store.get(queued["analysis_id"])
    job["execution"] = service.pipeline.prepare(job["request"], service.settings.data_dir / job["id"])
    job["snapshot"]["status"] = "analyzing"
    with service.store.transaction() as db:
        service.store.save(db, job)
    restarted = Service(service.settings)
    restarted.pipeline.prepare = Mock(side_effect=AssertionError("Must use the saved inputs"))
    assert restarted.run_once()
    assert restarted.get_analysis(queued["analysis_id"])["status"] == "awaiting_decision"
    assert calls.reviewer.call_count == 1


def test_crash_after_atomic_output_save_recovers_from_original_snapshot(backend):
    service, source, _ = backend
    snapshot = review(service)
    service.decide(snapshot["analysis_id"], decision(snapshot), "approve-1")
    job = service.store.get(snapshot["analysis_id"])
    output = service.pipeline.apply(job["execution"])  # Save, but simulate no completion transaction.
    restarted = Service(service.settings)
    assert restarted.run_once()
    assert restarted.get_analysis(snapshot["analysis_id"])["status"] == "applied"
    assert Document(output).paragraphs[0].text == "Wait 60 seconds"
    assert Document(source).paragraphs[0].text == "Wait 30 seconds"


def test_failures_are_visible_without_accidentally_accepting_or_writing(backend):
    service, _, calls = backend
    calls.reviewer.side_effect = RuntimeError("test failure")
    queued = service.start_analysis(request(), "start-1")[1]
    service.run_once()
    result = service.get_analysis(queued["analysis_id"])
    assert result["status"] == "failed" and result["error"]["code"] == "PIPELINE_FAILED"
    assert result["allowed_decisions"] == [] and not service.run_once()


def test_worker_lock_prevents_duplicate_processes_and_startup_recovers_queue(backend):
    service, _, _ = backend
    service.start_worker()
    other = Service(service.settings)
    try:
        with pytest.raises(RuntimeError, match="already uses"):
            other.start_worker()
    finally:
        service.stop_worker()
    other.start_worker(); other.stop_worker()


def test_backend_responses_match_the_frontend_openapi_contract(backend):
    from jsonschema import Draft202012Validator
    service, _, _ = backend
    contract_path = Path(__file__).resolve().parents[1] / "openapi.json"
    spec = json.loads(contract_path.read_text())
    sibling = contract_path.parent.parent / "github-app" / "backend.openapi.json"
    if sibling.exists():
        assert json.loads(sibling.read_text()) == spec
    def validate(value, name):
        schema = {"components": spec["components"], "$ref": f"#/components/schemas/{name}"}
        Draft202012Validator(schema).validate(value)
    with TestClient(create_app(service.settings, service, start_worker=False)) as client:
        validate(client.get("/v1/analyses/unknown").json(), "Error")
        validate(client.post("/v1/analyses", json={}, headers=AUTH).json(), "Error")
        queued = client.post("/v1/analyses", json=request(), headers=AUTH | {"Idempotency-Key": "schema-start"}).json()
        validate(request(), "StartAnalysis")
        validate(queued, "Analysis")
        service.run_once(); snapshot = service.get_analysis(queued["analysis_id"])
        validate(snapshot, "Analysis")
        response = client.post(f"/v1/analyses/{snapshot['analysis_id']}/decisions", json=decision(snapshot), headers=AUTH | {"Idempotency-Key": "schema-approve"}).json()
        validate(response, "DecisionAccepted")
        service.run_once(); validate(service.get_analysis(snapshot["analysis_id"]), "Analysis")
        assert client.get("/openapi.json").json()["components"]["securitySchemes"] == spec["components"]["securitySchemes"]


def test_api_uses_received_code_without_github_credentials_and_preserves_whitespace(backend):
    service, _, calls = backend
    body = request()
    file = body["code_snapshot"]["files"][0]
    file.update(patch="@@ -1 +1 @@\n-  timeout = 30\n+  timeout = 60\n", pr_head_content="  timeout = 60\n\n",
                merged_content="  timeout = 60\n\n# rebase context from other commits\n")
    with TestClient(create_app(service.settings, service, start_worker=False)) as client:
        response = client.post("/v1/analyses", json=body, headers=AUTH | {"Idempotency-Key": "supplied-code"})
        assert response.status_code == 202
    job = service.store.get(response.json()["analysis_id"])
    assert job["request"]["code_snapshot"]["files"][0] == file
    service.run_once()
    inputs = calls.reviewer.call_args.kwargs
    assert "  timeout = 30\n\n" in inputs["old_code"]
    assert "# rebase context" not in inputs["old_code"]
    assert "-  timeout = 30" in inputs["code_diff"] and "+  timeout = 60" in inputs["code_diff"]
    state = service.store.get(job["id"])["execution"]["state"]
    assert state["current_code"].endswith(file["merged_content"])


@pytest.mark.parametrize("kind", ["missing", "wrong_head", "wrong_merge", "duplicate", "removed_content"])
def test_http_rejects_absent_or_mismatched_source_snapshot(backend, kind):
    service, _, calls = backend
    body = request()
    if kind == "missing":
        del body["code_snapshot"]
    elif kind in ("wrong_head", "wrong_merge"):
        body["code_snapshot"]["head_sha" if kind == "wrong_head" else "merge_commit_sha"] = "d" * 40
    elif kind == "duplicate":
        body["code_snapshot"]["files"] *= 2
    else:
        body["code_snapshot"]["files"][0]["status"] = "removed"
    with TestClient(create_app(service.settings, service, start_worker=False)) as client:
        assert client.post("/v1/analyses", json=body, headers=AUTH | {"Idempotency-Key": kind}).status_code == 422
    assert calls.reviewer.call_count == 0 and not service.run_once()


def test_source_patch_corruption_fails_before_models(backend):
    service, _, calls = backend
    body = request(); body["code_snapshot"]["files"][0]["patch"] = "@@ -1 +1 @@\n-timeout = 30\n"
    queued = service.start_analysis(body, "truncated-input")[1]
    service.run_once()
    assert service.get_analysis(queued["analysis_id"])["status"] == "failed"
    assert calls.reviewer.call_count == 0


def test_oversized_source_snapshot_is_rejected_before_queueing(backend):
    service, _, calls = backend
    body = request(); body["code_snapshot"]["files"][0]["merged_content"] = "x" * (10 * 1024 * 1024)
    with TestClient(create_app(service.settings, service, start_worker=False)) as client:
        assert client.post("/v1/analyses", json=body, headers=AUTH | {"Idempotency-Key": "large"}).status_code == 422
    assert calls.reviewer.call_count == 0 and not service.run_once()


def test_javascript_collector_and_client_payload_drives_the_real_backend_flow(backend):
    import shutil
    import subprocess
    service, source, calls = backend
    app_dir = Path(__file__).resolve().parents[2] / "github-app"
    if not shutil.which("node") or not (app_dir / "src/clients.mjs").is_file():
        pytest.skip("Cross-project integration requires the sibling GitHub App and Node")
    script = r'''
import { collectSourceSnapshot } from './src/source.mjs';
import { BackendClient } from './src/clients.mjs';
const run = { id: 'node-to-python', attempt: 1, repository: 'owner/code', prNumber: 42,
  prUrl: 'https://github.com/owner/code/pull/42', baseSha: 'a'.repeat(40), headSha: 'b'.repeat(40),
  mergeCommitSha: 'c'.repeat(40), mergedAt: '2026-10-08T01:00:00Z', mergedBy: 'developer' };
const github = {
  async pr() { return { merged: true, merge_commit_sha: run.mergeCommitSha, changed_files: 1 }; },
  async pullFiles() { return [{ filename: 'worker.py', status: 'modified', additions: 1, deletions: 1,
    patch: '@@ -1 +1 @@\n-timeout = 30\n+timeout = 60\n' }]; },
  async fileText(context, path, sha) { return 'timeout = 60\n'; }
};
run.codeSnapshot = await collectSourceSnapshot(github, run);
const client = new BackendClient({});
client.request = async (method, path, body, key) => { process.stdout.write(JSON.stringify(body)); };
await client.start(run);
'''
    completed = subprocess.run(["node", "--input-type=module", "-e", script], cwd=app_dir,
                               capture_output=True, text=True, timeout=10, check=True)
    body = json.loads(completed.stdout)
    with TestClient(create_app(service.settings, service, start_worker=False)) as client:
        response = client.post("/v1/analyses", json=body, headers=AUTH | {"Idempotency-Key": "cross-project-start"})
        assert response.status_code == 202
        service.run_once()
        snapshot = client.get(f"/v1/analyses/{response.json()['analysis_id']}", headers=AUTH).json()
        assert snapshot["status"] == "awaiting_decision"
        assert "timeout = 30" in calls.reviewer.call_args.kwargs["old_code"]
        response = client.post(f"/v1/analyses/{snapshot['analysis_id']}/decisions", json=decision(snapshot),
                               headers=AUTH | {"Idempotency-Key": "cross-project-approve"})
        assert response.status_code == 202 and response.json()["accepted"]
        service.run_once()
        assert client.get(f"/v1/analyses/{snapshot['analysis_id']}", headers=AUTH).json()["status"] == "applied"
    assert Document(service.settings.data_dir / snapshot["analysis_id"] / "updated.docx").paragraphs[0].text == "Wait 60 seconds"
    assert Document(source).paragraphs[0].text == "Wait 30 seconds"


def test_invalid_or_oversized_proposal_fails_before_offering_approval(backend):
    service, _, calls = backend
    calls.proposer.return_value.changes[0].current_text = "Cannot find this text"
    queued = service.start_analysis(request(), "invalid-edit")[1]
    service.run_once()
    snapshot = service.get_analysis(queued["analysis_id"])
    assert snapshot["status"] == "failed" and not snapshot["proposals"]
    assert snapshot["proposal_version"] is None
    calls.proposer.return_value.changes[0].current_text = "Wait 30 seconds"
    calls.proposer.return_value.changes[0].proposed_text = "`" * 16000
    body = request(); body["attempt"] = 2
    queued = service.start_analysis(body, "oversized-fences")[1]; service.run_once()
    snapshot = service.get_analysis(queued["analysis_id"])
    assert snapshot["status"] == "failed" and not snapshot["proposals"]


def test_single_document_api_refuses_a_proposal_for_a_different_runbook(backend):
    service, _, calls = backend
    calls.proposer.return_value.changes[0].runbook = "Other Service.docx"
    queued = service.start_analysis(request(), "different-runbook")[1]
    service.run_once()
    snapshot = service.get_analysis(queued["analysis_id"])
    assert snapshot["status"] == "failed" and not snapshot["proposals"]


@pytest.mark.parametrize("module", ["src.agents.runbook_analyzer1", "src.agents.runbook_analyzer2", "src.agents.change_proposer"])
def test_agents_use_direct_content_without_retrieval(monkeypatch, module):
    import importlib
    from types import SimpleNamespace
    from test_workflow import code_review, proposal, review as runbook_review
    agent = importlib.import_module(module)
    monkeypatch.setattr(agent, "retrieve_for_queries", Mock(side_effect=AssertionError("No retrieval")))
    invoke = Mock(return_value=proposal() if module.endswith("change_proposer") else runbook_review())
    monkeypatch.setattr(agent, "llm", SimpleNamespace(with_structured_output=lambda schema: SimpleNamespace(invoke=invoke)))
    if module.endswith("change_proposer"):
        agent.create_change_proposal(runbook_review(), runbook_text="EXACT CONFIGURED DOCX TEXT")
    else:
        getattr(agent, module.rsplit(".", 1)[1])(code_review(), runbook_text="EXACT CONFIGURED DOCX TEXT")
    assert "EXACT CONFIGURED DOCX TEXT" in invoke.call_args.args[0]


def test_model_timeout_exposes_the_failed_step_without_request_details(backend):
    from httpx import Request
    from openai import APITimeoutError
    service, source, _ = backend
    original = source.read_bytes()
    error = APITimeoutError(request=Request("POST", "https://api.openai.com/v1/chat/completions",
                                          headers={"Authorization": "Bearer private-test-secret"}))
    error.add_note("During task with name 'change_proposal' and id 'test-task'")
    service.pipeline.analyze = Mock(side_effect=error)
    queued = service.start_analysis(request(), "timeout-start")[1]
    assert service.run_once()
    result = service.get_analysis(queued["analysis_id"])
    assert result["status"] == "failed"
    assert result["error"]["code"] == "MODEL_TIMEOUT"
    assert "generating proposed runbook changes" in result["error"]["message"]
    assert "private-test-secret" not in json.dumps(result)
    assert result["proposals"] == [] and result["allowed_decisions"] == []
    assert source.read_bytes() == original
    assert not service.run_once()
