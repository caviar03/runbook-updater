from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import pytest

from src.agents.implementor import implementor
from src.models.change_proposal import changeProposal, proposedChange


def proposal(*pairs):
    return changeProposal(readable_feedback="Review these edits", changes=[proposedChange(section="Recovery", current_text=old, proposed_text=new, reason="Code changed") for old, new in pairs])


def test_replacement_preserves_surrounding_text_run_formatting_and_original(tmp_path):
    source, output = tmp_path / "source.docx", tmp_path / "output.docx"
    document = Document()
    paragraph = document.add_paragraph(style="Heading 2")
    paragraph.add_run("Prefix ").italic = True
    paragraph.add_run("Wait 30 ").bold = True
    paragraph.add_run("seconds before retry. Suffix").underline = True
    document.add_paragraph("Unrelated content")
    document.save(source)
    before = source.read_bytes()
    implementor(str(source), proposal(("Wait 30 seconds before retry.", "Wait 60 seconds before retry.")), str(output))
    updated = Document(output)
    assert updated.paragraphs[0].text == "Prefix Wait 60 seconds before retry. Suffix"
    assert updated.paragraphs[0].runs[0].italic
    assert updated.paragraphs[0].runs[1].bold
    assert updated.paragraphs[0].runs[-1].underline
    assert updated.paragraphs[0].style.name == "Heading 2"
    assert updated.paragraphs[1].text == "Unrelated content"
    assert source.read_bytes() == before


def test_table_cell_and_nested_table_edits_are_found(tmp_path):
    source, output = tmp_path / "source.docx", tmp_path / "output.docx"
    document = Document()
    cell = document.add_table(rows=1, cols=1).cell(0, 0)
    cell.text = "Timeout: 30 seconds"
    nested = cell.add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "Retries: 2"
    document.save(source)
    implementor(str(source), proposal(("Timeout: 30 seconds", "Timeout: 60 seconds"), ("Retries: 2", "Retries: 3")), str(output))
    updated_cell = Document(output).tables[0].cell(0, 0)
    assert "Timeout: 60 seconds" in updated_cell.text
    assert updated_cell.tables[0].cell(0, 0).text == "Retries: 3"


@pytest.mark.parametrize("texts,old", [(["same", "same"], "same"), (["same same"], "same"), (["actual"], "missing"), (["actual"], "")])
def test_missing_ambiguous_or_empty_text_never_saves(tmp_path, texts, old):
    source, output = tmp_path / "source.docx", tmp_path / "output.docx"
    document = Document()
    for text in texts:
        document.add_paragraph(text)
    document.save(source)
    output.write_bytes(b"existing output")
    with pytest.raises(ValueError):
        implementor(str(source), proposal((old, "replacement")), str(output))
    assert output.read_bytes() == b"existing output"
    assert not list(tmp_path.glob("tmp*.docx"))


def test_failure_in_a_later_change_does_not_partially_overwrite_source(tmp_path):
    source = tmp_path / "source.docx"
    document = Document(); document.add_paragraph("first"); document.save(source)
    before = source.read_bytes()
    with pytest.raises(ValueError, match="found 0"):
        implementor(str(source), proposal(("first", "changed"), ("missing", "new")), str(source))
    assert source.read_bytes() == before


def test_merged_table_cells_are_not_counted_as_duplicate_matches(tmp_path):
    source, output = tmp_path / "source.docx", tmp_path / "output.docx"
    document = Document()
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Only once"
    document.save(source)
    implementor(str(source), proposal(("Only once", "Updated once")), str(output))
    assert Document(output).tables[0].cell(0, 0).text == "Updated once"


def test_empty_approval_cannot_write_a_document(tmp_path):
    with pytest.raises(ValueError, match="human-approved"):
        implementor("unused.docx", proposal(), str(tmp_path / "output.docx"))


def test_edit_cannot_destroy_a_field_in_the_same_run(tmp_path):
    source, output = tmp_path / "source.docx", tmp_path / "output.docx"
    document = Document()
    run = document.add_paragraph().add_run("Timeout: 30 seconds")
    field = OxmlElement("w:fldChar")
    field.set(qn("w:fldCharType"), "begin")
    run._r.append(field)
    document.save(source)
    before = source.read_bytes()
    with pytest.raises(ValueError, match="overlaps an image, field"):
        implementor(str(source), proposal(("30 seconds", "60 seconds")), str(output))
    assert source.read_bytes() == before and not output.exists()
