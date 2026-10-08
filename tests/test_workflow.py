import importlib
from types import SimpleNamespace
from unittest.mock import Mock

from docx import Document
import pytest

from src.graph import nodes
from src.graph.workflow import build_workflow
from src.main import make_initial_state
from src.models.code_review import CodeChangeReview, CodeChange
from src.models.runbook_review import runbookReview, runbookIssue
from src.models.change_proposal import changeProposal, proposedChange
from src.models.orchestrator_decision import orchestratorDecision


def code_review():
    return CodeChangeReview(existing_discrepancies=[], code_changes=[CodeChange(description="Timeout changed", affected_components=["worker"], operational_impact="Wait longer", runbook_relevant=True)])


def review(update=True):
    return runbookReview(update_required=update, issues=[runbookIssue(section="Recovery", issue="Timeout outdated", evidence_from_code="60 seconds", recommended_change="Use 60 seconds", priority="high")] if update else [])


def proposal(empty=False):
    return changeProposal(readable_feedback="Update timeout", changes=[] if empty else [proposedChange(section="Recovery", current_text="Wait 30 seconds", proposed_text="Wait 60 seconds", reason="Timeout changed")])


def setup_workflow(monkeypatch, tmp_path, update=True):
    source, output = tmp_path / "selected.docx", tmp_path / "updated.docx"
    document = Document(); document.add_paragraph("Wait 30 seconds"); document.save(source)
    state = make_initial_state({"old_code": "OLD", "current_code": "NEW", "code_diff": "DIFF", "changed_files": []}, source, output)
    reviewer = Mock(return_value=code_review())
    analyzer1, analyzer2 = Mock(return_value=review(update)), Mock(return_value=review(update))
    judge = Mock(return_value=review(update))
    orchestrator = Mock(return_value=orchestratorDecision(proceed=True, reasoning="Supported", feedback=""))
    proposer = Mock(return_value=proposal())
    monkeypatch.setattr(nodes, "code_reviewer", reviewer)
    monkeypatch.setattr(nodes, "runbook_analyzer1", analyzer1)
    monkeypatch.setattr(nodes, "runbook_analyzer2", analyzer2)
    monkeypatch.setattr(nodes, "runbook_decision_reviewer", judge)
    monkeypatch.setattr(nodes, "orchestrator_review", orchestrator)
    monkeypatch.setattr(nodes, "create_change_proposal", proposer)
    return state, source, output, SimpleNamespace(reviewer=reviewer, analyzer1=analyzer1, analyzer2=analyzer2, judge=judge, orchestrator=orchestrator, proposer=proposer)


@pytest.mark.parametrize("module", ["src.agents.runbook_analyzer1", "src.agents.runbook_analyzer2"])
def test_analyzer_prompt_contains_code_retrieved_evidence_and_rework_feedback(monkeypatch, module):
    agent = importlib.import_module(module)
    retrieval = Mock(return_value=[SimpleNamespace(id="chunk", payload={"source": "worker.docx", "headings": ["Recovery"], "text": "RUNBOOK EVIDENCE: timeout is 30 seconds"})])
    invoke = Mock(return_value=review())
    monkeypatch.setattr(agent, "retrieve_for_queries", retrieval)
    monkeypatch.setattr(agent, "llm", SimpleNamespace(with_structured_output=lambda schema: SimpleNamespace(invoke=invoke)))
    getattr(agent, module.rsplit(".", 1)[1])(code_review(), feedback="REWORK: verify timeout evidence")
    prompt = invoke.call_args.args[0]
    assert "Timeout changed" in prompt
    assert "RUNBOOK EVIDENCE: timeout is 30 seconds" in prompt
    assert "REWORK: verify timeout evidence" in prompt


def test_approved_graph_uses_correct_inputs_and_selected_document(monkeypatch, tmp_path):
    state, source, output, calls = setup_workflow(monkeypatch, tmp_path)
    monkeypatch.setattr("builtins.input", lambda prompt: "yes")
    result = build_workflow().invoke(state)
    assert result["workflow_status"] == "applied"
    assert result["updated_runbook_path"] == str(output)
    assert Document(output).paragraphs[0].text == "Wait 60 seconds"
    assert Document(source).paragraphs[0].text == "Wait 30 seconds"
    calls.reviewer.assert_called_once_with(old_code="OLD", code_diff="DIFF")
    assert calls.judge.call_count == 1


def test_no_update_exits_without_prompt_or_document_write(monkeypatch, tmp_path):
    state, _, output, calls = setup_workflow(monkeypatch, tmp_path, update=False)
    monkeypatch.setattr("builtins.input", Mock(side_effect=AssertionError("No approval should be requested")))
    result = build_workflow().invoke(state)
    assert result["workflow_status"] == "no_update_needed"
    assert calls.proposer.call_count == 0 and calls.orchestrator.call_count == 0
    assert not output.exists()


def test_empty_proposal_does_not_ask_for_approval(monkeypatch, tmp_path):
    state, _, output, calls = setup_workflow(monkeypatch, tmp_path)
    calls.proposer.return_value = proposal(empty=True)
    monkeypatch.setattr("builtins.input", Mock(side_effect=AssertionError("Empty proposal must not be approved")))
    assert build_workflow().invoke(state)["workflow_status"] == "no_update_needed"
    assert not output.exists()


def test_denial_performs_no_write(monkeypatch, tmp_path):
    state, source, output, _ = setup_workflow(monkeypatch, tmp_path)
    before = source.read_bytes()
    monkeypatch.setattr("builtins.input", lambda prompt: "no")
    assert build_workflow().invoke(state)["workflow_status"] == "rejected"
    assert source.read_bytes() == before and not output.exists()


def test_orchestrator_rework_reruns_both_analyzers_with_feedback(monkeypatch, tmp_path):
    state, _, _, calls = setup_workflow(monkeypatch, tmp_path)
    calls.orchestrator.side_effect = [orchestratorDecision(proceed=False, reasoning="Incomplete", feedback="Check the timeout"), orchestratorDecision(proceed=True, reasoning="Supported", feedback="")]
    monkeypatch.setattr("builtins.input", lambda prompt: "no")
    result = build_workflow().invoke(state)
    assert result["analysis_attempts"] == 2
    assert calls.judge.call_count == calls.analyzer1.call_count == calls.analyzer2.call_count == 2
    assert calls.analyzer1.call_args_list[-1].kwargs["feedback"] == "Check the timeout"
    assert calls.analyzer2.call_args_list[-1].kwargs["feedback"] == "Check the timeout"


def test_rework_attempts_stop_without_recursion_error_or_write(monkeypatch, tmp_path):
    state, _, output, calls = setup_workflow(monkeypatch, tmp_path)
    calls.orchestrator.return_value = orchestratorDecision(proceed=False, reasoning="Unsupported", feedback="Insufficient evidence")
    state["max_analysis_attempts"] = 2
    result = build_workflow().invoke(state)
    assert result["workflow_status"] == "failed"
    assert result["analysis_attempts"] == calls.judge.call_count == 2
    assert not output.exists()


def test_human_edit_feedback_regenerates_instead_of_ending_as_rejected(monkeypatch, tmp_path):
    state, _, output, calls = setup_workflow(monkeypatch, tmp_path)
    answers = iter(["edit", "Clarify the waiting period", "yes"])
    monkeypatch.setattr("builtins.input", lambda prompt: next(answers))
    result = build_workflow().invoke(state)
    assert result["workflow_status"] == "applied" and output.exists()
    assert calls.proposer.call_count == 2
    assert calls.analyzer1.call_args_list[-1].kwargs["feedback"] == "Clarify the waiting period"


def test_cli_state_defaults_to_updated_copy_and_keeps_defined_state_keys():
    state = make_initial_state({"old_code": "old", "current_code": "new", "code_diff": "diff", "changed_files": []}, "selected.docx")
    assert state["output_path"] == "selected.updated.docx"
    assert state["code_diff"] == "diff"
    from src.models.workflow import workflowState
    assert set(state) <= set(workflowState.__annotations__)
