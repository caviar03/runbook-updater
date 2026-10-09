import json
import sys
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from types import ModuleType, SimpleNamespace
from unittest.mock import Mock

import pytest
from docx import Document
from fastapi.testclient import TestClient

from src.api.app import create_app
from src.api.service import Service
from src.api.settings import Settings
from src.models.change_proposal import changeProposal, proposedChange
from src.retrieval.runtime import current_scope
from test_api import request, decision, TOKEN, AUTH
from test_workflow import setup_workflow, review as agent_review


@pytest.fixture
def backend_v2(monkeypatch, tmp_path):
    _, _, _, calls = setup_workflow(monkeypatch,tmp_path)
    paths = []
    for name in ['Orders.docx','Payments.docx','Search.docx']:
        path=tmp_path/name
        document=Document(); document.add_paragraph('Recovery',style='Heading 1'); document.add_paragraph('Wait 30 seconds'); document.save(path)
        paths.append(path)
    calls.proposer.return_value=changeProposal(readable_feedback='Orders and payments need changes',changes=[
        proposedChange(runbook=path.name,section='Recovery',current_text='Wait 30 seconds',proposed_text='Wait 60 seconds',reason='Timeout changed') for path in paths[:2]])
    indexer=Mock()
    module=ModuleType('src.retrieval.index_runbooks'); module.index_runbooks=indexer
    monkeypatch.setitem(sys.modules,'src.retrieval.index_runbooks',module)
    monkeypatch.setattr('src.retrieval.runtime.client_for',lambda *args: SimpleNamespace(count=lambda **kwargs: SimpleNamespace(count=3)))
    settings=Settings(TOKEN,'owner/code',paths[0],tmp_path/'data',tuple(paths))
    return Service(settings),paths,calls,indexer


def analyzed(service):
    queued=service.start_analysis(request(),'start')[1]
    assert service.run_once()
    result=service.get_analysis(queued['analysis_id'])
    assert result['status']=='awaiting_decision',result
    return result


def choose(service,snapshot,index,action='approve',key=None):
    proposal=snapshot['proposals'][index]
    return service.decide_document(snapshot['analysis_id'],proposal['id'],decision(snapshot,action),key or proposal['id']+action)


def test_api_calls_existing_indexer_and_graph_without_full_document_override(backend_v2):
    service,paths,calls,indexer=backend_v2
    snapshot=analyzed(service)
    job=service.store.get(snapshot['analysis_id'])
    assert job['execution']['format']==2
    assert 'runbook_text' not in job['execution']['state']
    assert calls.analyzer1.call_args.kwargs.get('runbook_text') is None
    assert calls.proposer.call_args.kwargs.get('runbook_text') is None
    assert [Path(path).name for path in indexer.call_args.args[0]]==[path.name for path in paths]
    assert all('sources' in path for path in indexer.call_args.args[0])
    assert len(snapshot['proposals'])==2
    assert {proposal['target']['path'] for proposal in snapshot['proposals']}=={'Orders.docx','Payments.docx'}
    assert snapshot['decision_scope']=='document'
    assert not (Path(job['execution']['state']['output_dir'])).exists()
    assert current_scope() is None


def test_independent_decisions_wait_for_every_document_and_save_only_approved(backend_v2):
    service,paths,_,_=backend_v2
    originals=[path.read_bytes() for path in paths]
    snapshot=analyzed(service)
    accepted=choose(service,snapshot,0)[1]['analysis']
    assert accepted['status']=='awaiting_decision'
    assert not service.run_once()
    assert accepted['proposals'][0]['review_status']=='approved'
    with pytest.raises(Exception,match='document decision endpoint'):
        service.decide(snapshot['analysis_id'],decision(snapshot),'bundle')
    assert choose(service,snapshot,1,'deny')[1]['analysis']['status']=='applying'
    assert service.run_once()
    final=service.get_analysis(snapshot['analysis_id'])
    assert final['status']=='applied'
    approved=final['proposals'][0]
    output=Path(approved['result']['output_path'])
    assert Document(output).paragraphs[1].text=='Wait 60 seconds'
    assert not (output.parent/'Payments.docx').exists()
    assert not (output.parent/'Search.docx').exists()
    assert [path.read_bytes() for path in paths]==originals
    assert final['proposals'][1]['review_status']=='denied'


def test_all_denied_never_queues_a_save_and_no_update_requires_ack(backend_v2):
    service,_,calls,_=backend_v2
    snapshot=analyzed(service)
    choose(service,snapshot,0,'deny'); choose(service,snapshot,1,'deny')
    assert not service.run_once()
    assert service.get_analysis(snapshot['analysis_id'])['status']=='rejected'
    calls.analyzer1.return_value=calls.analyzer2.return_value=calls.judge.return_value=agent_review(False)
    body=request();body['attempt']=2
    queued=service.start_analysis(body,'no-update')[1]
    assert service.run_once()
    snapshot=service.get_analysis(queued['analysis_id'])
    assert snapshot['status']=='no_update_needed'
    assert snapshot['allowed_decisions']==['acknowledge']
    assert service.decide(snapshot['analysis_id'],decision(snapshot,'acknowledge'),'ack')[1]['accepted']
    assert not service.run_once()


def test_duplicate_and_competing_document_decisions_and_restart(backend_v2):
    service,_,_,_=backend_v2
    snapshot=analyzed(service)
    with ThreadPoolExecutor(max_workers=2) as executor:
        results=list(executor.map(lambda _:choose(service,snapshot,0,key='same-key'),range(2)))
    assert results[0]==results[1]
    restarted=Service(service.settings)
    with pytest.raises(Exception,match='not awaiting'):
        choose(restarted,snapshot,0,'deny')
    choose(restarted,snapshot,1)
    assert restarted.run_once()
    assert restarted.get_analysis(snapshot['analysis_id'])['status']=='applied'


def test_source_changed_before_approval_or_batch_save_is_stale(backend_v2):
    service,paths,_,_=backend_v2
    snapshot=analyzed(service)
    choose(service,snapshot,0)
    paths[0].write_bytes(paths[0].read_bytes()+b'changed')
    choose(service,snapshot,1)
    assert service.run_once()
    final=service.get_analysis(snapshot['analysis_id'])
    assert final['status']=='stale'
    assert all(proposal['save_status']=='not_started' for proposal in final['proposals'])
    assert final['proposals'][0]['review_status']=='approved'


def test_failed_save_preserves_other_receipts_and_does_not_claim_success(backend_v2,monkeypatch):
    service,_,_,_=backend_v2
    snapshot=analyzed(service);choose(service,snapshot,0);choose(service,snapshot,1)
    original=service.pipeline.apply_document
    def apply(execution,name,changes):
        if name=='Payments.docx':raise OSError('disk write failed')
        return original(execution,name,changes)
    monkeypatch.setattr(service.pipeline,'apply_document',apply)
    assert service.run_once()
    final=service.get_analysis(snapshot['analysis_id'])
    assert final['status']=='failed'
    assert final['proposals'][0]['save_status']=='saved'
    assert final['proposals'][1]['save_status']=='failed'
    assert Path(final['proposals'][0]['result']['output_path']).is_file()


def test_retrieval_failure_is_visible_and_no_direct_text_fallback(backend_v2):
    service,_,calls,indexer=backend_v2
    indexer.side_effect=RuntimeError('index unavailable')
    queued=service.start_analysis(request(),'start')[1]
    assert service.run_once()
    final=service.get_analysis(queued['analysis_id'])
    assert final['status']=='failed' and final['error']['code']=='RETRIEVAL_FAILED'
    assert not calls.reviewer.called
    assert final['proposals']==[]


def test_parallel_graph_analyzers_inherit_their_own_retrieval_collection(backend_v2,monkeypatch):
    from src.graph import nodes
    service,_,_,_=backend_v2
    scopes=[]
    def analyzer(**kwargs):
        assert 'runbook_text' not in kwargs
        scopes.append(current_scope()[1])
        return agent_review()
    monkeypatch.setattr(nodes,'runbook_analyzer1',analyzer)
    monkeypatch.setattr(nodes,'runbook_analyzer2',analyzer)
    snapshot=analyzed(service)
    expected=service.store.get(snapshot['analysis_id'])['execution']['retrieval']['collection']
    assert scopes==[expected,expected]


def test_new_document_endpoint_auth_and_contract(backend_v2):
    from jsonschema import Draft202012Validator
    service,_,_,_=backend_v2
    snapshot=analyzed(service)
    schema=json.loads((Path(__file__).parents[1]/'openapi.json').read_text())
    validator=Draft202012Validator({'$ref':'#/components/schemas/Analysis','components':schema['components']})
    validator.validate(snapshot)
    url=f"/v1/analyses/{snapshot['analysis_id']}/proposals/{snapshot['proposals'][0]['id']}/decisions"
    with TestClient(create_app(service.settings,service,start_worker=False)) as client:
        assert client.post(url,json=decision(snapshot)).status_code==401
        result=client.post(url,json=decision(snapshot),headers=AUTH|{'Idempotency-Key':'http-doc'})
        assert result.status_code==202
        validator.validate(result.json()['analysis'])


def test_multi_runbook_settings_validate_paths_and_preserve_cli_defaults(tmp_path,monkeypatch):
    paths=[]
    for name in ['one.docx','two.docx']:
        path=tmp_path/name;Document().save(path);paths.append(path)
    monkeypatch.setenv('RUNBOOK_PATHS',json.dumps([str(path) for path in paths]))
    monkeypatch.setenv('BACKEND_TOKEN',TOKEN);monkeypatch.setenv('SOURCE_REPOSITORY','owner/code')
    settings=Settings.from_env()
    assert settings.paths==tuple(paths)
    assert settings.retrieval_embedded
    monkeypatch.setenv('RUNBOOK_PATHS','[]')
    with pytest.raises(ValueError):Settings.from_env()


def test_restart_skips_a_completed_document_save_and_resumes_the_rest(backend_v2,monkeypatch):
    service,_,_,_=backend_v2
    snapshot=analyzed(service);choose(service,snapshot,0);choose(service,snapshot,1)
    original=service.pipeline.apply_document
    class ProcessExit(BaseException):pass
    def apply(execution,name,changes):
        if name=='Payments.docx':raise ProcessExit()
        return original(execution,name,changes)
    monkeypatch.setattr(service.pipeline,'apply_document',apply)
    with pytest.raises(ProcessExit):service.run_once()
    restarted=Service(service.settings)
    resumed=Mock(wraps=restarted.pipeline.apply_document)
    monkeypatch.setattr(restarted.pipeline,'apply_document',resumed)
    assert restarted.run_once()
    assert [call.args[1] for call in resumed.call_args_list]==['Payments.docx']
    assert restarted.get_analysis(snapshot['analysis_id'])['status']=='applied'


def test_unknown_document_is_rejected_and_bad_versions_cannot_accept_a_decision(backend_v2):
    service,_,calls,_=backend_v2
    snapshot=analyzed(service)
    body=decision(snapshot);body['proposal_version']='old'
    with pytest.raises(Exception,match='does not match'):
        service.decide_document(snapshot['analysis_id'],snapshot['proposals'][0]['id'],body,'stale')
    assert service.get_analysis(snapshot['analysis_id'])['proposals'][0]['review_status']=='pending'
    calls.proposer.return_value.changes[0].runbook='Unknown.docx'
    request2=request();request2['attempt']=2
    queued=service.start_analysis(request2,'unknown')[1]
    assert service.run_once()
    assert service.get_analysis(queued['analysis_id'])['status']=='failed'
